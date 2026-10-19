
import csv
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

WORKBOOK_PATH = (
    ROOT
    / "data/raw/neighbourhoods/neighbourhood-profiles-2021-158-model.xlsx"
)

COVERAGE_CSV = (
    ROOT
    / "data/processed/neighbourhoods/neighbourhood_washroom_coverage.csv"
)

COVERAGE_GEOJSON = (
    ROOT
    / "data/processed/neighbourhoods/neighbourhood_washroom_coverage.geojson"
)

OUTPUT_DIR = ROOT / "data/processed/neighbourhoods"

OUTPUT_CSV = OUTPUT_DIR / "neighbourhood_population_coverage.csv"
OUTPUT_GEOJSON = OUTPUT_DIR / "neighbourhood_population_coverage.geojson"


def normalize_code(value):
    """Normalize neighbourhood identifiers to integer strings."""
    return str(int(value))


def load_population():
    wb = load_workbook(
        WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        ws = wb["hd2021_census_profile"]

        names = list(
            next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        )

        codes = list(
            next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        )

        population_row = list(
            next(ws.iter_rows(min_row=4, max_row=4, values_only=True))
        )

        expected_indicator = (
            "Total - Age groups of the population - 25% sample data"
        )

        if population_row[0].strip() != expected_indicator:
            raise ValueError(
                "Population indicator on row 4 has changed"
            )

        populations = {}

        for column in range(1, len(codes)):
            code = codes[column]
            name = names[column]
            population = population_row[column]

            if code is None:
                continue

            key = normalize_code(code)

            if key in populations:
                raise ValueError(
                    f"Duplicate neighbourhood code: {key}"
                )

            if (
                not isinstance(population, (int, float))
                or population <= 0
            ):
                raise ValueError(
                    f"Invalid population for {name}: {population}"
                )

            populations[key] = {
                "name": str(name),
                "population": int(population),
            }

        return populations

    finally:
        wb.close()


def main():
    print("=" * 65)
    print("CIVICACCESS POPULATION-NORMALIZED COVERAGE")
    print("=" * 65)

    populations = load_population()

    with open(COVERAGE_CSV, encoding="utf-8") as file:
        coverage_rows = list(csv.DictReader(file))

    with open(COVERAGE_GEOJSON, encoding="utf-8") as file:
        geojson = json.load(file)

    print(f"Population records: {len(populations)}")
    print(f"Coverage records: {len(coverage_rows)}")
    print(f"Neighbourhood polygons: {len(geojson['features'])}")

    coverage_codes = {
        normalize_code(row["neighbourhood_code"])
        for row in coverage_rows
    }

    population_codes = set(populations)

    missing_population = coverage_codes - population_codes
    missing_coverage = population_codes - coverage_codes

    if missing_population or missing_coverage:
        raise ValueError(
            "Neighbourhood identifier mismatch.\n"
            f"Missing population: {sorted(missing_population)}\n"
            f"Missing coverage: {sorted(missing_coverage)}"
        )

    enriched_rows = []
    rows_by_code = {}

    for row in coverage_rows:
        code = normalize_code(row["neighbourhood_code"])

        population = populations[code]["population"]
        candidate_sites = int(row["candidate_site_count"])
        accessible_reported = int(
            row["accessibility_reported_site_count"]
        )

        sites_per_10000 = (
            candidate_sites / population * 10000
        )

        reported_accessibility_per_10000 = (
            accessible_reported / population * 10000
        )

        enriched = {
            **row,
            "population_2021_estimate": population,
            "sites_per_10000_residents": round(
                sites_per_10000, 3
            ),
            "accessibility_reported_sites_per_10000": round(
                reported_accessibility_per_10000, 3
            ),
            "population_source_indicator": (
                "Total - Age groups of the population "
                "- 25% sample data"
            ),
        }

        enriched_rows.append(enriched)
        rows_by_code[code] = enriched

    enriched_rows.sort(
        key=lambda row: (
            float(row["sites_per_10000_residents"]),
            row["neighbourhood_name"],
        )
    )

    enriched_features = []

    for feature in geojson["features"]:
        code = normalize_code(
            feature["properties"]["AREA_SHORT_CODE"]
        )

        enriched_features.append({
            "type": "Feature",
            "properties": {
                **feature["properties"],
                **rows_by_code[code],
            },
            "geometry": feature["geometry"],
        })

    with open(
        OUTPUT_CSV,
        "w",
        newline="",
        encoding="utf-8",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=list(enriched_rows[0].keys()),
        )
        writer.writeheader()
        writer.writerows(enriched_rows)

    with open(OUTPUT_GEOJSON, "w", encoding="utf-8") as file:
        json.dump(
            {
                "type": "FeatureCollection",
                "features": enriched_features,
            },
            file,
        )

    print("\nVALIDATION")
    print("-" * 65)

    total_sites = sum(
        int(row["candidate_site_count"])
        for row in enriched_rows
    )

    zero_site_count = sum(
        int(row["candidate_site_count"]) == 0
        for row in enriched_rows
    )

    print(f"Total assigned sites: {total_sites}")
    print(f"Neighbourhoods with zero sites: {zero_site_count}")
    print(f"Population matched: {len(enriched_rows)}")

    assert len(enriched_rows) == 158
    assert len(enriched_features) == 158
    assert total_sites == 366
    assert zero_site_count == 26

    print("Validation passed.")

    print("\nTEN LOWEST SITE RATES")
    print("-" * 65)

    for row in enriched_rows[:10]:
        print(
            f"{row['neighbourhood_name']}: "
            f"{row['candidate_site_count']} sites, "
            f"population {row['population_2021_estimate']}, "
            f"{row['sites_per_10000_residents']} "
            f"sites per 10,000 residents"
        )

    print("\nOUTPUT FILES")
    print(OUTPUT_CSV)
    print(OUTPUT_GEOJSON)


if __name__ == "__main__":
    main()
